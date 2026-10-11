
# ============================================================
# SISTEMA DE SEGUIMIENTO DE VACUNACIÓN INFANTIL
# ARCHIVO: backend/services/sisvacu_service.py
#
# FUNCIÓN:
# Leer, interpretar y validar los datos del Excel SISVACU.
#
# OPERACIONES:
# 1. Abrir el archivo Excel.
# 2. Identificar las columnas.
# 3. Leer los pacientes.
# 4. Validar nombres y documentos.
# 5. Interpretar edades.
# 6. Identificar vacunas aplicadas.
# 7. Detectar errores y advertencias.
# 8. Preparar los datos para su importación.
#
# IMPORTANTE:
# Este archivo NO guarda información en SQLite.
# ============================================================


# ============================================================
# 1. IMPORTAR LIBRERÍAS
# ============================================================

# io:
# Permite leer el archivo Excel desde la memoria.

import io

# re:
# Permite validar textos utilizando expresiones regulares.

import re

# unicodedata:
# Permite eliminar acentos y normalizar textos.

import unicodedata

# openpyxl:
# Permite leer archivos Excel .xlsx.

from openpyxl import load_workbook


# ============================================================
# 2. COLUMNAS PRINCIPALES DEL SISVACU
# ============================================================

# Estas columnas contienen información del paciente.
#
# Las demás columnas representan vacunas o dosis.

COLUMNAS_BASICAS = [
    "Servicio",
    "Edad",
    "Documento",
    "Persona"
]


# ============================================================
# 3. NORMALIZAR TEXTO
# ============================================================

def normalizar_texto(texto):
    """
    Convierte el texto a mayúsculas, elimina espacios
    innecesarios y elimina acentos.

    Ejemplo:

    José Ramón Pérez
        -> JOSE RAMON PEREZ
    """

    # Si el valor está vacío, devolvemos texto vacío.

    if texto is None:
        return ""

    # Convertimos el valor a texto y mayúsculas.

    texto = str(texto).strip().upper()

    # Eliminamos espacios repetidos.

    texto = " ".join(texto.split())

    # Eliminamos los acentos.

    texto = "".join(
        caracter
        for caracter in unicodedata.normalize("NFD", texto)
        if unicodedata.category(caracter) != "Mn"
    )

    return texto


# ============================================================
# 4. INTERPRETAR EDAD
# ============================================================

def interpretar_edad(valor):
    """
    Extrae el número de años del valor de SISVACU.

    Ejemplos:

    "5 año/s" -> 5
    "2 año/s" -> 2
    "0 año/s" -> 0

    IMPORTANTE:
    No calcula la edad exacta en meses.
    """

    if valor is None:
        return None

    texto = str(valor).strip()

    # Buscar el primer número del texto.

    coincidencia = re.search(
        r"\d+",
        texto
    )

    if not coincidencia:
        return None

    return int(coincidencia.group())


# ============================================================
# 5. NORMALIZAR DOCUMENTO
# ============================================================

def normalizar_documento(valor):
    """
    Prepara el documento para guardarlo.

    Ejemplos:

    8841439.0 -> 8841439
    jscg07012021 -> JSCG07012021
    isaias -> ISAIAS

    No intentamos interpretar los códigos temporales.
    """

    if valor is None:
        return None

    # Excel puede representar una cédula como decimal.

    if isinstance(valor, float):
        if valor.is_integer():
            valor = int(valor)

    documento = str(valor).strip().upper()

    # Corregimos documentos numéricos terminados en .0.

    if documento.endswith(".0"):
        posible_numero = documento[:-2]

        if posible_numero.isdigit():
            documento = posible_numero

    # Quitamos espacios internos.

    documento = documento.replace(" ", "")

    if not documento:
        return None

    return documento


# ============================================================
# 6. CLASIFICAR DOCUMENTO
# ============================================================

def clasificar_documento(documento):
    """
    Clasifica el identificador del paciente.

    Posibles resultados:

    cedula
    temporal_sisvacu
    sin_documento
    invalido

    Los identificadores temporales no se interpretan
    como documentos de padres o tutores.
    """

    # --------------------------------------------------------
    # PACIENTE SIN DOCUMENTO
    # --------------------------------------------------------

    if not documento:
        return "sin_documento"

    # --------------------------------------------------------
    # DOCUMENTO NUMÉRICO
    # --------------------------------------------------------

    if documento.isdigit():

        # Cédula numérica con longitud esperada.
        if 5 <= len(documento) <= 8:
            return "cedula"

        return "invalido"

    # --------------------------------------------------------
    # IDENTIFICADOR TEMPORAL SISVACU
    # --------------------------------------------------------

    # Ejemplos:
    #
    # JSCG07012021
    # MLGO04092026
    # ISAIAS

    if re.fullmatch(r"[A-Z0-9_-]+", documento):
        return "temporal_sisvacu"

    # --------------------------------------------------------
    # DOCUMENTO NO RECONOCIDO
    # --------------------------------------------------------

    return "invalido"


# ============================================================
# 7. VALIDAR NOMBRE
# ============================================================

def validar_nombre(nombre):
    """
    Comprueba que el nombre contenga caracteres permitidos.

    Acepta:
    - Letras
    - Espacios
    - Acentos
    - Ñ
    - Guiones
    - Apóstrofes
    - Puntos

    Ejemplo válido:
    R.N MILAGROS GONZALEZ GONZALEZ
    """

    if not nombre:
        return False

    patron = (
        r"^[A-Za-zÁÉÍÓÚÜÑ"
        r"áéíóúüñ.' -]+$"
    )

    return bool(
        re.fullmatch(
            patron,
            str(nombre).strip()
        )
    )


# ============================================================
# 8. SEPARAR NOMBRES Y APELLIDOS
# ============================================================

def separar_nombre(nombre_completo):
    """
    SISVACU proporciona el nombre completo en una
    sola columna llamada Persona.

    Como no existen columnas independientes, realizamos
    una separación provisional.

    IMPORTANTE:
    Esta separación debe poder corregirse manualmente.
    """

    partes = str(nombre_completo).strip().split()

    # Sin nombre.

    if len(partes) == 0:
        return ("SIN NOMBRE", "SIN APELLIDO")

    # Solo una palabra.

    if len(partes) == 1:
        return (partes[0], "SIN APELLIDO")

    # Dos palabras.

    if len(partes) == 2:
        return (partes[0], partes[1])

    # Tres o más palabras:
    # dividimos aproximadamente por la mitad.

    mitad = len(partes) // 2

    nombres = " ".join(partes[:mitad])
    apellidos = " ".join(partes[mitad:])

    return (nombres, apellidos)


# ============================================================
# 9. LEER Y VALIDAR EXCEL SISVACU
# ============================================================

def leer_sisvacu(contenido_archivo):
    """
    Función principal del lector SISVACU.

    Recibe:
        contenido_archivo: bytes del archivo Excel.

    Devuelve:
        Un diccionario con los pacientes, vacunas,
        errores, advertencias y estadísticas.
    """

    # --------------------------------------------------------
    # LISTAS DE CONTROL
    # --------------------------------------------------------

    errores_generales = []

    filas_resultado = []

    # Detecta identificadores repetidos dentro del Excel.

    identificadores_en_archivo = set()

    # --------------------------------------------------------
    # ABRIR ARCHIVO EXCEL
    # --------------------------------------------------------

    try:

        archivo_memoria = io.BytesIO(
            contenido_archivo
        )

        workbook = load_workbook(
            archivo_memoria,
            data_only=True
        )

    except Exception as error:

        return {
            "valido": False,
            "errores": [
                f"No se pudo abrir el archivo Excel: {error}"
            ],
            "total_filas": 0,
            "filas_ok": 0,
            "filas_advertencia": 0,
            "filas_error": 0,
            "columnas_vacunas": [],
            "filas": []
        }

    # --------------------------------------------------------
    # SELECCIONAR HOJA
    # --------------------------------------------------------

    hoja = workbook.active

    # --------------------------------------------------------
    # LEER ENCABEZADOS
    # --------------------------------------------------------

    # En el Excel utilizado:
    #
    # Fila 1: título.
    # Fila 2: filtros.
    # Fila 3: encabezados.
    # Fila 4 en adelante: pacientes.

    encabezados = []

    for celda in hoja[3]:

        if celda.value is None:
            encabezados.append(None)

        else:
            encabezados.append(
                str(celda.value).strip()
            )

    # --------------------------------------------------------
    # COMPROBAR COLUMNAS OBLIGATORIAS
    # --------------------------------------------------------

    for columna in COLUMNAS_BASICAS:

        if columna not in encabezados:

            errores_generales.append(
                f"Falta la columna obligatoria: {columna}"
            )

    if errores_generales:

        return {
            "valido": False,
            "errores": errores_generales,
            "total_filas": 0,
            "filas_ok": 0,
            "filas_advertencia": 0,
            "filas_error": 0,
            "columnas_vacunas": [],
            "filas": []
        }

    # --------------------------------------------------------
    # IDENTIFICAR COLUMNAS DE VACUNAS
    # --------------------------------------------------------

    columnas_vacunas = [
        columna
        for columna in encabezados
        if columna and columna not in COLUMNAS_BASICAS
    ]

    # --------------------------------------------------------
    # IDENTIFICAR POSICIONES DE COLUMNAS
    # --------------------------------------------------------

    indice_servicio = encabezados.index("Servicio")
    indice_edad = encabezados.index("Edad")
    indice_documento = encabezados.index("Documento")
    indice_persona = encabezados.index("Persona")

    # --------------------------------------------------------
    # RECORRER FILAS DEL EXCEL
    # --------------------------------------------------------

    for numero_fila in range(4, hoja.max_row + 1):

        valores = [
            hoja.cell(
                row=numero_fila,
                column=i + 1
            ).value
            for i in range(len(encabezados))
        ]

        # Ignorar filas completamente vacías.

        if not any(
            valor is not None
            for valor in valores
        ):
            continue

        errores = []
        advertencias = []

        # ----------------------------------------------------
        # OBTENER INFORMACIÓN DEL PACIENTE
        # ----------------------------------------------------

        servicio = valores[indice_servicio]

        edad_original = valores[indice_edad]

        documento = normalizar_documento(
            valores[indice_documento]
        )

        persona = valores[indice_persona]

        # ----------------------------------------------------
        # VALIDAR NOMBRE
        # ----------------------------------------------------

        if not persona:

            errores.append(
                "El nombre del paciente está vacío."
            )

        elif not validar_nombre(persona):

            errores.append(
                "El nombre del paciente contiene caracteres no válidos."
            )

        # ----------------------------------------------------
        # VALIDAR DOCUMENTO
        # ----------------------------------------------------

        tipo_documento = clasificar_documento(
            documento
        )

        if tipo_documento == "sin_documento":

            advertencias.append(
                "Paciente sin documento. Debe completarse posteriormente."
            )

        elif tipo_documento == "invalido":

            errores.append(
                f"El documento tiene un formato no reconocido: {documento}"
            )

        elif tipo_documento == "temporal_sisvacu":

            advertencias.append(
                "El paciente utiliza un identificador temporal SISVACU."
            )

        # ----------------------------------------------------
        # DETECTAR IDENTIFICADORES REPETIDOS
        # ----------------------------------------------------

        if documento:

            clave_documento = (
                tipo_documento,
                documento
            )

            if clave_documento in identificadores_en_archivo:

                errores.append(
                    "El identificador está repetido dentro del archivo."
                )

            else:

                identificadores_en_archivo.add(
                    clave_documento
                )

        # ----------------------------------------------------
        # VALIDAR EDAD
        # ----------------------------------------------------

        edad = interpretar_edad(
            edad_original
        )

        if edad is None:

            errores.append(
                "No se pudo interpretar la edad."
            )

        elif edad < 0:

            errores.append(
                "La edad no puede ser negativa."
            )

        elif edad > 6:

            advertencias.append(
                "La edad está fuera del rango esperado de 0 a 6 años."
            )

        # ----------------------------------------------------
        # LEER VACUNAS APLICADAS
        # ----------------------------------------------------

        dosis_aplicadas = []

        for indice, encabezado in enumerate(encabezados):

            # Ignorar columnas sin nombre.

            if not encabezado:
                continue

            # Ignorar columnas principales.

            if encabezado in COLUMNAS_BASICAS:
                continue

            valor = valores[indice]

            # Una celda vacía no registra una dosis.
            # No significa necesariamente que esté pendiente.

            if (
                valor is None
                or str(valor).strip() == ""
            ):
                continue

            # El valor 1 representa una dosis registrada.

            if (
                valor == 1
                or str(valor).strip() in ["1", "1.0"]
            ):

                dosis_aplicadas.append(
                    encabezado
                )

            else:

                errores.append(
                    f"La dosis {encabezado} tiene un valor inválido: {valor}"
                )

        # ----------------------------------------------------
        # PACIENTE SIN DOSIS MARCADAS
        # ----------------------------------------------------

        if not dosis_aplicadas:

            advertencias.append(
                "El paciente no tiene dosis marcadas en el archivo."
            )

        # ----------------------------------------------------
        # PREPARAR NOMBRES Y APELLIDOS
        # ----------------------------------------------------

        nombres = None
        apellidos = None

        if persona:

            nombres, apellidos = separar_nombre(
                persona
            )

        # ----------------------------------------------------
        # DETERMINAR ESTADO DE LA FILA
        # ----------------------------------------------------

        if errores:
            estado = "error"

        elif advertencias:
            estado = "advertencia"

        else:
            estado = "ok"

        # ----------------------------------------------------
        # GUARDAR RESULTADO DE LA VALIDACIÓN
        # ----------------------------------------------------

        filas_resultado.append({
            "fila": numero_fila,
            "estado": estado,
            "servicio": servicio,
            "edad_original": edad_original,
            "edad": edad,
            "documento": documento,
            "tipo_documento": tipo_documento,
            "persona": persona,
            "nombres": nombres,
            "apellidos": apellidos,
            "dosis_aplicadas": dosis_aplicadas,
            "errores": errores,
            "advertencias": advertencias
        })

    # --------------------------------------------------------
    # CALCULAR ESTADÍSTICAS
    # --------------------------------------------------------

    filas_ok = sum(
        1
        for fila in filas_resultado
        if fila["estado"] == "ok"
    )

    filas_advertencia = sum(
        1
        for fila in filas_resultado
        if fila["estado"] == "advertencia"
    )

    filas_error = sum(
        1
        for fila in filas_resultado
        if fila["estado"] == "error"
    )

    # --------------------------------------------------------
    # DEVOLVER RESULTADO FINAL
    # --------------------------------------------------------

    return {
        "valido": True,
        "errores": [],
        "total_filas": len(filas_resultado),
        "filas_ok": filas_ok,
        "filas_advertencia": filas_advertencia,
        "filas_error": filas_error,
        "columnas_vacunas": columnas_vacunas,
        "filas": filas_resultado
    }
